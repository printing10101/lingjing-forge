"""Document identity, shortening, retry and rollback contracts against real ChromaDB."""
from __future__ import annotations

from unittest.mock import patch

import pytest

from app.rag.document_importer import DocumentImportService
from app.rag.entity_index import EntityIndex
from app.rag.knowledge_base import KnowledgeStore, CollectionProxy
from app.rag.vector_store import VectorStore

pytestmark = pytest.mark.unit


class Embeddings:
    def embed_batch(self, chunks):
        return [[1., 0., 0.] for _ in chunks]


@pytest.fixture
def library(tmp_path, monkeypatch):
    from app.rag import document_importer
    store = KnowledgeStore.__new__(KnowledgeStore)
    store._vs = VectorStore(persist_directory=str(tmp_path / "chroma"))
    store._entity_index = EntityIndex()
    store._next_id = 1
    store._emb = Embeddings()
    class Library:
        collection = CollectionProxy(store)
    monkeypatch.setattr(document_importer, "get_embedding_service", lambda: store._emb)
    monkeypatch.setattr(document_importer, "_chunk_text", lambda text, *_: text.splitlines())
    importer = DocumentImportService(Library())
    yield importer, store
    store._vs.close()


def test_same_named_documents_are_separate_and_shorter_revision_removes_tail(library, tmp_path):
    importer, store = library
    first = tmp_path / "factory-a" / "manual.txt"
    second = tmp_path / "factory-b" / "manual.txt"
    first.parent.mkdir()
    second.parent.mkdir()
    first.write_text("加工45钢工件\n旧切削力数据\n旧转速数据", encoding="utf-8")
    second.write_text("铝合金参数", encoding="utf-8")
    a = importer.import_document(str(first))
    b = importer.import_document(str(second))
    assert a["success"] and b["success"]
    assert a["document_id"] != b["document_id"]
    assert store.count() == 4
    first.write_text("新版45钢参数", encoding="utf-8")
    updated = importer.import_document(str(first))
    assert updated["success"] and updated["document_revision"] != a["document_revision"]
    assert store.count() == 2
    assert store._entity_index.get_chunks(["切削力"]) == []
    documents = store._vs.get(limit=10)["documents"]
    assert sorted(documents) == sorted(["新版45钢参数", "铝合金参数"])
    assert importer.import_document(str(first))["success"]
    assert store.count() == 2


def test_failed_replacement_restores_previous_chunks(library, tmp_path):
    importer, store = library
    source = tmp_path / "manual.txt"
    source.write_text("旧45钢资料\n旧转速", encoding="utf-8")
    assert importer.import_document(str(source))["success"]
    previous = store._vs.get(limit=10)
    source.write_text("新铝合金资料", encoding="utf-8")
    collection = store._vs._collection
    delete = collection.delete
    calls = 0
    def fail_once(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 1:
            raise RuntimeError("injected stale chunk deletion failure")
        return delete(*args, **kwargs)
    with patch.object(collection, "delete", side_effect=fail_once):
        assert importer.import_document(str(source))["success"] is False
    after = store._vs.get(limit=10)
    assert sorted(zip(after["ids"], after["documents"])) == sorted(zip(previous["ids"], previous["documents"]))
    assert store._entity_index.get_chunks(["铝合金"]) == []


def test_excel_legacy_format_has_actionable_error(tmp_path):
    from app.rag.excel_parser import parse_excel
    legacy = tmp_path / "legacy.xls"
    legacy.write_bytes(b"not used: extension is rejected before parsing")
    result = parse_excel(legacy)
    assert result["status"] == "error" and "转存" in result["error"]


def test_excel_merged_cells_and_uncached_formulas(tmp_path):
    import openpyxl
    from app.rag.excel_parser import parse_excel
    source = tmp_path / "table.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["材料", "转速"])
    sheet.append(["45钢", 1000])
    sheet.append([None, 1200])
    sheet.merge_cells("A2:A3")
    book.save(source)
    result = parse_excel(source)
    assert result["status"] == "success"
    assert result["rows"] == [["45钢", "1000"], ["45钢", "1200"]]
    sheet["B3"] = "=B2+200"
    book.save(source)
    book.close()
    result = parse_excel(source)
    assert result["status"] == "error" and "重新计算" in result["error"]


def test_word_tables_and_excel_rows_reach_same_document_store(library, tmp_path):
    import openpyxl
    from docx import Document
    importer, store = library
    word = tmp_path / "process.docx"
    document = Document()
    document.add_paragraph("45钢工艺规范")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "材料"
    table.cell(0, 1).text = "转速(rpm)"
    table.cell(1, 0).text = "45钢"
    table.cell(1, 1).text = "1200"
    document.save(word)
    workbook = openpyxl.Workbook()
    workbook.active.title = "刀具参数"
    workbook.active.append(["材料", "刀具直径(mm)"])
    workbook.active.append(["45钢", 10])
    sheet = tmp_path / "tools.xlsx"
    workbook.save(sheet)
    workbook.close()
    assert importer.import_document(str(word))["success"]
    assert importer.import_document(str(sheet))["success"]
    content = "\n".join(store._vs.get(limit=100)["documents"])
    assert "1200" in content and "刀具参数" in content and "刀具直径(mm)" in content
    ids = store._entity_index.get_chunks(["45钢"])
    metadata = store._vs.get(ids=ids, limit=100)["metadatas"]
    assert {item["file_type"] for item in metadata} == {".docx", ".xlsx"}


def test_invalid_word_is_a_failed_import_without_store_changes(library, tmp_path):
    importer, store = library
    source = tmp_path / "invalid.docx"
    source.write_bytes(b"invalid word package")
    result = importer.import_document(str(source))
    assert result["success"] is False
    assert store.count() == 0
